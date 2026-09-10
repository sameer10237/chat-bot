from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).parent / "SDIT_AI_Campus_Assistant_Presentation.pptx"

NAVY = RGBColor(13, 27, 42)
NAVY_2 = RGBColor(20, 42, 63)
CYAN = RGBColor(52, 211, 219)
ORANGE = RGBColor(247, 148, 30)
WHITE = RGBColor(247, 250, 252)
MUTED = RGBColor(177, 196, 209)
GREEN = RGBColor(81, 190, 133)
RED = RGBColor(235, 104, 104)


def add_text(slide, text, x, y, w, h, size=20, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def circle(slide, x, y, d, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    return shape


def line(slide, x1, y1, x2, y2, color=CYAN, width=2):
    shape = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def base_slide(prs, title, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    rect(slide, 0, 0, 13.333, 0.12, CYAN)
    add_text(slide, title, 0.65, 0.42, 11.8, 0.55, 28, WHITE, True)
    add_text(slide, f"SDIT AI Campus Assistant  |  {number:02d}", 0.68, 7.18, 5.0, 0.22, 9, MUTED)
    add_text(slide, "5-minute project presentation", 9.25, 7.18, 3.4, 0.22, 9, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def bullet_list(slide, items, x, y, w, h, size=18, accent=CYAN):
    for i, item in enumerate(items):
        yy = y + i * 0.58
        circle(slide, x, yy + 0.13, 0.12, accent)
        add_text(slide, item, x + 0.28, yy, w - 0.28, 0.43, size, WHITE)


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 0, 0, 13.333, 0.16, CYAN)
    rect(slide, 8.65, 0.16, 4.68, 7.34, NAVY_2)
    circle(slide, 9.55, 1.15, 2.5, CYAN)
    circle(slide, 10.17, 1.77, 1.26, NAVY)
    rect(slide, 9.28, 4.28, 3.35, 0.12, ORANGE)
    add_text(slide, "SDIT AI CAMPUS\nASSISTANT", 0.75, 1.28, 7.5, 1.5, 34, WHITE, True)
    add_text(slide, "A grounded chatbot for student questions", 0.8, 3.18, 6.6, 0.5, 22, CYAN)
    add_text(slide, "Shree Devi Institute of Technology, Mangalore", 0.8, 4.0, 6.4, 0.38, 17, MUTED)
    add_text(slide, "Built with Python, FastAPI, RAG and NVIDIA NIM", 0.8, 5.35, 6.6, 0.35, 15, WHITE)
    add_text(slide, "Project presentation | 5 minutes", 0.8, 6.65, 4.5, 0.3, 12, MUTED)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)

    slide = base_slide(prs, "1. The problem we wanted to solve", 2)
    add_text(slide, "Students repeatedly ask the same campus questions, but information is spread across many web pages.", 0.75, 1.25, 11.7, 0.5, 20, CYAN, True)
    bullet_list(slide, [
        "Finding admissions, departments, facilities and placement information takes time.",
        "A normal chatbot may guess when it does not know the answer.",
        "The system must be useful even when an external AI API is unavailable.",
        "Answers should be traceable to official SDIT information."
    ], 1.0, 2.15, 11.0, 3.0, 19, ORANGE)
    rect(slide, 0.78, 5.75, 11.8, 0.7, NAVY_2, True)
    add_text(slide, "Goal: provide quick, reliable and source-grounded answers in a simple chat interface.", 1.05, 5.86, 11.2, 0.45, 20, WHITE, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, "2. Our solution", 3)
    add_text(slide, "A two-layer assistant: retrieve trusted SDIT content first, then answer from that content.", 0.75, 1.15, 11.6, 0.45, 19, MUTED)
    boxes = [(0.75, "Student\nquestion", ORANGE), (3.1, "FastAPI\nweb app", CYAN), (5.45, "RAG\nretrieval", GREEN), (7.8, "NVIDIA NIM\n(optional)", ORANGE), (10.15, "Grounded\nanswer", CYAN)]
    for i, (x, label, color) in enumerate(boxes):
        rect(slide, x, 2.45, 1.85, 1.25, NAVY_2, True, color)
        add_text(slide, label, x + 0.12, 2.65, 1.61, 0.85, 17, WHITE, True, PP_ALIGN.CENTER)
        if i < len(boxes) - 1:
            line(slide, x + 1.85, 3.08, x + 2.28, 3.08, color)
    rect(slide, 1.45, 4.65, 10.3, 1.15, NAVY_2, True)
    add_text(slide, "If NVIDIA is configured: natural-language generation\nIf not configured or unavailable: local RAG fallback", 1.7, 4.83, 9.8, 0.75, 21, WHITE, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, "3. Knowledge and RAG pipeline", 4)
    add_text(slide, "Official SDIT website content is converted into structured data before the chatbot is used.", 0.75, 1.1, 11.7, 0.45, 19, MUTED)
    steps = [
        (0.8, "1", "Scrape", "Collect SDIT pages"),
        (3.35, "2", "Clean", "Organize sections"),
        (5.9, "3", "Chunk", "Create searchable passages"),
        (8.45, "4", "Match", "Find relevant Q&A + chunks"),
        (11.0, "5", "Answer", "Return with sources"),
    ]
    for x, num, head, desc in steps:
        circle(slide, x, 2.35, 0.72, CYAN)
        add_text(slide, num, x, 2.46, 0.72, 0.35, 20, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, head, x - 0.35, 3.35, 1.45, 0.35, 17, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, desc, x - 0.75, 3.78, 2.25, 0.6, 13, MUTED, align=PP_ALIGN.CENTER)
    for x in [1.55, 4.1, 6.65, 9.2]:
        line(slide, x, 2.71, x + 1.05, 2.71, ORANGE)
    rect(slide, 1.0, 5.2, 11.2, 0.8, NAVY_2, True)
    add_text(slide, "Current dataset: 45 pages  |  296 knowledge chunks  |  240 multi-form Q&A pairs", 1.15, 5.38, 10.9, 0.4, 20, WHITE, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, "4. What happens when a user asks a question?", 5)
    add_text(slide, "Example: 'What is the eligibility criteria for Computer Science Engineering?'", 0.75, 1.1, 11.7, 0.45, 20, CYAN, True)
    flow = [
        (0.8, "Input", "Validate question\nand conversation"),
        (3.35, "Retrieve", "Search Q&A pairs\nand RAG chunks"),
        (5.9, "Ground", "Build context only\nfrom SDIT records"),
        (8.45, "Generate", "NVIDIA model or\nlocal fallback"),
        (11.0, "Respond", "Answer + source\nlinks"),
    ]
    for i, (x, head, desc) in enumerate(flow):
        rect(slide, x, 2.25, 1.65, 1.45, NAVY_2, True, CYAN if i != 3 else ORANGE)
        add_text(slide, head, x + 0.08, 2.5, 1.49, 0.3, 17, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.08, 2.92, 1.49, 0.58, 12, MUTED, align=PP_ALIGN.CENTER)
        if i < len(flow) - 1:
            line(slide, x + 1.65, 2.98, x + 2.08, 2.98, CYAN)
    rect(slide, 1.1, 4.65, 11.0, 1.05, NAVY_2, True)
    add_text(slide, "Trust rule: when the knowledge base does not contain the answer, the bot says so instead of inventing information.", 1.35, 4.88, 10.5, 0.55, 19, GREEN, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, "5. Main features", 6)
    features = [
        ("Grounded answers", "Uses official SDIT records and shows source links.", CYAN),
        ("Dual operating mode", "Works with NVIDIA NIM or local RAG fallback.", ORANGE),
        ("Conversation aware", "Understands short follow-up questions using recent context.", GREEN),
        ("Safe failure", "Handles missing keys, provider errors and off-topic requests.", RED),
        ("Student-friendly UI", "Suggestions, markdown, copy, retry and loading states.", CYAN),
        ("CLI included", "The same assistant can also run in the terminal.", ORANGE),
    ]
    for i, (head, desc, color) in enumerate(features):
        col = i % 2
        row = i // 2
        x = 0.9 + col * 6.05
        y = 1.35 + row * 1.55
        rect(slide, x, y, 5.45, 1.1, NAVY_2, True, color)
        circle(slide, x + 0.25, y + 0.36, 0.28, color)
        add_text(slide, head, x + 0.72, y + 0.18, 4.4, 0.3, 17, WHITE, True)
        add_text(slide, desc, x + 0.72, y + 0.55, 4.4, 0.35, 13, MUTED)

    slide = base_slide(prs, "6. Technology stack and project structure", 7)
    stack = [(0.9, "Frontend", "HTML + CSS + JavaScript", CYAN), (4.65, "Backend", "Python + FastAPI", ORANGE), (8.4, "AI layer", "NVIDIA NIM + local RAG", GREEN)]
    for x, head, desc, color in stack:
        rect(slide, x, 1.45, 3.15, 1.35, NAVY_2, True, color)
        add_text(slide, head, x + 0.15, 1.67, 2.85, 0.3, 19, color, True, PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.15, 2.08, 2.85, 0.35, 14, WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, "Important files", 0.9, 3.55, 2.0, 0.35, 19, WHITE, True)
    bullet_list(slide, [
        "app.py - web server and chat API endpoints",
        "rag_engine.py - search, scoring and context retrieval",
        "nvidia_client.py - model calls and fallback behavior",
        "data/ - cleaned SDIT knowledge, chunks and Q&A pairs"
    ], 1.0, 4.05, 10.8, 2.1, 17, CYAN)

    slide = base_slide(prs, "7. Demo plan and conclusion", 8)
    add_text(slide, "A simple 5-minute live explanation", 0.8, 1.1, 5.5, 0.4, 21, CYAN, True)
    demo = [
        "1. Open the chatbot web page.",
        "2. Ask an admissions or department question.",
        "3. Show the answer and its source links.",
        "4. Ask an unsupported question to show safe fallback.",
        "5. Explain that the system still works without the NVIDIA key."
    ]
    bullet_list(slide, demo, 1.0, 1.8, 6.0, 3.1, 17, ORANGE)
    rect(slide, 7.65, 1.35, 4.55, 3.9, NAVY_2, True, GREEN)
    add_text(slide, "Final message", 8.0, 1.75, 3.85, 0.35, 20, GREEN, True, PP_ALIGN.CENTER)
    add_text(slide, "This project combines\nweb scraping, retrieval,\nlarge language models\nand responsible fallback\ninto one useful campus assistant.", 8.05, 2.35, 3.75, 2.15, 22, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "Thank you", 8.4, 5.8, 3.1, 0.45, 25, ORANGE, True, PP_ALIGN.CENTER)

    prs.save(OUT)
    print(f"Created {OUT}")


if __name__ == "__main__":
    build_deck()